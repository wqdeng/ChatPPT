import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches

from logger import LOG  # 引入日志模块

# 生成 PowerPoint 演示文稿
def generate_presentation(template_path: str, output_path: str):
    # 检查模板文件是否存在
    if not os.path.exists(template_path):
        LOG.error(f"模板文件 '{template_path}' 不存在。")  # 记录错误日志
        raise FileNotFoundError(f"模板文件 '{template_path}' 不存在。")

    prs = Presentation(template_path)  # 加载 PowerPoint 模板
    prs.core_properties.title = "Slide Demo"  # 设置 PowerPoint 的核心标题

    # =============== 首页 ===============
    home_page = prs.slides.add_slide(prs.slide_layouts[0])
    home_page.shapes.title.text = "Slide Demo"

    # =============== 文本 ===============
    content_page = prs.slides.add_slide(prs.slide_layouts[6])
    content_page.shapes.title.text = 'Effective delivery techniques'
    content_page.placeholders[2].text = "This is a powerful tool in public speaking. It involves varying pitch, tone, and volume to convey emotion, emphasize points, and maintain interest."
    p1 = content_page.placeholders[2].text_frame.add_paragraph()
    p1.text = "• Pitch variation"
    p2 = content_page.placeholders[2].text_frame.add_paragraph()
    p2.text = "• Tone inflection"
    p3 = content_page.placeholders[2].text_frame.add_paragraph()
    p3.text = "• Volume control"

    content_page.placeholders[4].text = "Effective body language enhances your message, making it more impactful and memorable."
    p4 = content_page.placeholders[4].text_frame.add_paragraph()
    p4.text = "• Meaningful eye contact"
    p5 = content_page.placeholders[4].text_frame.add_paragraph()
    p5.text = "• Purposeful gestures"
    p6 = content_page.placeholders[4].text_frame.add_paragraph()
    p6.text = "• Maintain good posture"
    p7 = content_page.placeholders[4].text_frame.add_paragraph()
    p7.text = "• Control your expressions"

    # =============== 图片 ===============
    picture_page = prs.slides.add_slide(prs.slide_layouts[8])
    picture_page.shapes.title.text = "Speaking impact"

    image_full_path = os.path.join(os.path.dirname(__file__), "../images/demo1.png") # 构建图片的绝对路径
    for shape in picture_page.placeholders:
        if shape.placeholder_format.type == 2:
            picture_page.placeholders[shape.placeholder_format.idx].text = "Your ability to communicate effectively will leave a lasting impact on your audience\nEffectively communicating involves not only delivering a message but also resonating with the experiences, values, and emotions of those listening"
        if shape.placeholder_format.type == 18:  # 18 表示图片占位符
            shape.insert_picture(image_full_path)

    # =============== 表格 ===============
    table_page = prs.slides.add_slide(prs.slide_layouts[11])
    table_page.shapes.title.text = "Speaking engagement metrics"
    left = Inches(1.40)   # 表格左边距
    top = Inches(2.30)  # 表格上边距
    width = Inches(11)  # 表格宽度
    height = Inches(4.5) # 表格高度

    draft_table = table_page.shapes.add_table(6, 4, left, top, width, height).table
    data_title = ["Impact factor", "Measurement", "Target", "Achieved"]
    for i, cel in enumerate(data_title):
        cell = draft_table.cell(0, i)
        cell.text = cel

        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(40, 72, 255)
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER # 设置文本水平居中
        cell.text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE # 设置文本垂直居中

    data_content = [
        ['Audience interaction', 'Percentage (%)', '85', '88'],
        ['Knowledge retention', 'Percentage (%)', '75', '80'],
        ['Post-presentation surveys', 'Average rating', '4.2', '4.5'],
        ['Referral rate', 'Percentage (%)', '10', '12'],
        ['Collaboration opportunities', '# of opportunities', '8', '10'],
    ]
    for i, row in enumerate(data_content, 1):
        for j, cel in enumerate(row):
            cell = draft_table.cell(i, j)
            cell.text = cel

            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER # 设置文本水平居中
            cell.text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE # 设置文本垂直居中

            cell.fill.solid()
            if i % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            else:
                cell.fill.fore_color.rgb = RGBColor(235, 235, 255)

    # =============== 图表 ===============
    chart_page = prs.slides.add_slide(prs.slide_layouts[13])
    chart_page.shapes.title.text = "Chart Demo"
    # 创建图表数据
    chart_data = CategoryChartData()
    chart_data.categories = ['类别 1', '类别 2', '类别 3']
    chart_data.add_series('系列 1', (1.5, 2.7, 3.2))
    chart_data.add_series('系列 2', (2.1, 3.8, 4.0))

    # 添加柱状图
    x, y, cx, cy = Inches(2), Inches(2), Inches(6), Inches(4.5)
    chart = chart_page.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart

    # 保存生成的 PowerPoint 文件
    prs.save(output_path)
    LOG.info(f"演示文稿已保存到 '{output_path}'")

if __name__ == "__main__":
    template_path = os.path.join(os.path.dirname(__file__), "../templates/GalaxyPresentation.pptx")
    print(template_path)
    output_path = os.path.join(os.path.dirname(__file__), "../outputs/Slide_Demo.pptx")
    generate_presentation(template_path, output_path)
